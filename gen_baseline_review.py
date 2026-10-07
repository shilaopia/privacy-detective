# -*- coding: utf-8 -*-
"""合并 baseline_v7.json + review_json/*.json,生成横向对比 Excel。

结构:主 sheet 行=58 条款(列:一级维度/二级维度/核心条款/义务细则数 + 每目标2列),
汇总 sheet = 11 目标 × 4 判定计数 + 合规率。
样式沿用 make_report.py 约定:GREEN/YELLOW/RED/GRAY = 合规/存疑/不合规/未提及。
"""
import json
import os
import glob
import openpyxl
import config
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = config.BASE
OUT = os.path.join(BASE, "privacy_excels", "制度标尺v7-横向对比.xlsx")

# 目标列序(与名单文件一致:直接型C/B + 嵌入型)
TARGETS = [
    ("千问_ToC", "千问\nTo-C"), ("千问_ToB", "千问\nTo-B"),
    ("淘宝_ToC", "淘宝"),
    ("元宝_ToC", "元宝\nTo-C"), ("元宝_ToB", "元宝\nTo-B"),
    ("微信_ToC", "微信"),
    ("豆包_ToC", "豆包\nTo-C"), ("豆包_ToB", "豆包\nTo-B"),
    ("抖音_ToC", "抖音"),
    ("DeepSeek_ToC", "DeepSeek\nTo-C"), ("DeepSeek_ToB", "DeepSeek\nTo-B"),
]

FILL = {
    "合规": PatternFill("solid", fgColor="C6EFCE"),
    "存疑": PatternFill("solid", fgColor="FFEB9C"),
    "不合规": PatternFill("solid", fgColor="FFC7CE"),
    "未提及": PatternFill("solid", fgColor="D9D9D9"),
}
DIM_FILL = PatternFill("solid", fgColor="DEEAF1")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    baseline = json.load(open(os.path.join(BASE, "baseline_v7.json"), encoding="utf-8"))
    reviews = {}
    for key, _ in TARGETS:
        p = os.path.join(BASE, "review_json", key + ".json")
        if not os.path.exists(p):
            raise SystemExit(f"缺少审查结果: {p}")
        data = json.load(open(p, encoding="utf-8"))
        # {条款名: {"评估":..., "合规说明":...}}
        reviews[key] = data

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "制度标尺审查"

    # 表头两行
    fixed = ["一级维度", "二级维度", "核心对标条款", "义务\n细则数"]
    ws.append(fixed + [t for _, name in TARGETS for t in (name, "")])
    ws.append(["", "", "", ""] + ["评估", "合规说明"] * len(TARGETS))
    for c in range(1, 5):
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(1, c).font = HEAD_FONT
        ws.cell(1, c).alignment = CENTER
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for i, (_, name) in enumerate(TARGETS):
        c1 = 5 + i * 2
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c1 + 1)
        ws.cell(1, c1).fill = HEAD_FILL
        ws.cell(1, c1).font = HEAD_FONT
        ws.cell(1, c1).alignment = CENTER
        ws.cell(1, c1 + 1).fill = HEAD_FILL
        for r in (1, 2):
            ws.cell(r, c1).fill = HEAD_FILL
            ws.cell(r, c1 + 1).fill = HEAD_FILL
        ws.cell(2, c1).font = HEAD_FONT
        ws.cell(2, c1 + 1).font = HEAD_FONT
        ws.cell(2, c1).alignment = CENTER
        ws.cell(2, c1 + 1).alignment = CENTER
    for c in range(1, 5):
        ws.cell(2, c).fill = HEAD_FILL

    # 数据行
    r = 3
    dim1_rows = {}
    for cl in baseline:
        dim1_rows.setdefault(cl["dim1"], []).append(r)
        ws.cell(r, 1, cl["dim1"])
        ws.cell(r, 2, cl["dim2"] or "—")
        ws.cell(r, 3, cl["clause"])
        ws.cell(r, 4, cl["detail_count"])
        for i, (key, _) in enumerate(TARGETS):
            cell_data = reviews[key].get(cl["clause"], {})
            verdict = cell_data.get("评估", "未提及")
            note = cell_data.get("合规说明", "")
            cv = ws.cell(r, 5 + i * 2, verdict)
            cv.fill = FILL.get(verdict, FILL["未提及"])
            cv.alignment = CENTER
            ws.cell(r, 6 + i * 2, note).alignment = WRAP
        r += 1

    # 一级维度合并 + 着色
    for dim, rows in dim1_rows.items():
        if len(rows) > 1:
            ws.merge_cells(start_row=rows[0], start_column=1, end_row=rows[-1], end_column=1)
        for rr in rows:
            ws.cell(rr, 1).fill = DIM_FILL
            ws.cell(rr, 1).alignment = CENTER

    # 样式收尾
    for row in ws.iter_rows(min_row=3, max_row=r - 1, max_col=4):
        for cell in row:
            cell.border = THIN
            if cell.column in (2, 3):
                cell.alignment = WRAP
            elif cell.column == 4:
                cell.alignment = CENTER
    for row in ws.iter_rows(min_row=3, max_row=r - 1, min_col=5, max_col=4 + 2 * len(TARGETS)):
        for cell in row:
            cell.border = THIN
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=4 + 2 * len(TARGETS)):
        for cell in row:
            cell.border = THIN

    ws.freeze_panes = "E3"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 7
    for i in range(len(TARGETS)):
        ws.column_dimensions[get_column_letter(5 + i * 2)].width = 6
        ws.column_dimensions[get_column_letter(6 + i * 2)].width = 42
    ws.row_dimensions[1].height = 30
    for rr in range(3, r):
        ws.row_dimensions[rr].height = 60

    # 汇总 sheet
    ws2 = wb.create_sheet("汇总")
    ws2.append(["目标", "合规", "存疑", "不合规", "未提及", "合规率(合规/58)"])
    for cell in ws2[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
    for key, name in TARGETS:
        counts = {v: 0 for v in FILL}
        for cl in baseline:
            v = reviews[key].get(cl["clause"], {}).get("评估", "未提及")
            counts[v if v in counts else "未提及"] += 1
        ws2.append([name.replace("\n", ""), counts["合规"], counts["存疑"],
                    counts["不合规"], counts["未提及"], f"{counts['合规']/len(baseline):.0%}"])
    ws2.column_dimensions["A"].width = 16
    for col in "BCDEF":
        ws2.column_dimensions[col].width = 12

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    main()
