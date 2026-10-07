# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import config

GREEN  = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED    = PatternFill("solid", fgColor="FFC7CE")
GRAY   = PatternFill("solid", fgColor="D9D9D9")
HEADER = PatternFill("solid", fgColor="2F5496")
SUBHDR = PatternFill("solid", fgColor="BDD7EE")
BLUE1  = PatternFill("solid", fgColor="DEEAF1")
BLUE2  = PatternFill("solid", fgColor="E2EFDA")

thin   = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

FILL_MAP  = {"合规": GREEN, "存疑": YELLOW, "不合规": RED, "未提及": GRAY}
RISK_FILL = {"高": RED, "中": YELLOW, "低": GREEN}


def _sc(ws, row, col, value, fill=None, font=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:  c.fill  = fill
    if font:  c.font  = font
    if align: c.alignment = align
    c.border = border
    return c


def _verdict_fill(verdict: str):
    for key, fill in FILL_MAP.items():
        if key in verdict:
            return fill
    return None


def _risk_fill(risk: str):
    return RISK_FILL.get(risk)


def generate_platform_report(
    platform_name: str,
    policy_date: str,
    framework_note: str,
    rows: list,
    risks: list,
    output_path: str,
):
    """
    Generate a compliance audit Excel report for one platform.

    rows: list of 8-element lists:
        [stage, dimension, item, quote, location, verdict, explanation, risk_level]
        verdict examples: "合规" / "存疑" / "不合规" / "未提及"
        risk_level: "高" / "中" / "低" / ""

    risks: list of 4-element lists:
        [priority_label, description, location, suggestion]
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{platform_name}隐私政策审查"

    # Title
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"{platform_name} 隐私政策合规审查报告（基于生成式AI平台实证研究框架）"
    c.fill = HEADER
    c.font = Font(bold=True, color="FFFFFF", size=13)
    c.alignment = center
    c.border = border

    # Subtitle
    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = f"平台：{platform_name}  |  政策版本：{policy_date}  |  审查框架：{framework_note}"
    c.fill = SUBHDR
    c.font = Font(size=9, italic=True, color="1F3864")
    c.alignment = center
    c.border = border

    # Column headers
    headers = ["阶段", "审查维度", "具体审查项", "原文摘录（关键句）", "出处", "合规判断", "说明", "风险等级"]
    for i, h in enumerate(headers, 1):
        _sc(ws, 3, i, h,
            fill=PatternFill("solid", fgColor="1F4E79"),
            font=Font(bold=True, color="FFFFFF", size=10),
            align=center)

    # Data rows
    for i, row in enumerate(rows, 4):
        stage   = row[0]
        verdict = row[5]
        risk    = row[7]
        for j, val in enumerate(row, 1):
            f = None
            if j == 6:
                f = _verdict_fill(verdict)
            elif j == 8:
                f = _risk_fill(risk)
            elif j == 1:
                f = BLUE1 if "训练" in stage else BLUE2
            _sc(ws, i, j, val,
                fill=f,
                font=Font(size=9),
                align=left if j > 2 else center)

    # Column widths / row heights
    for i, w in enumerate([14, 16, 22, 44, 20, 12, 38, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows():
        ws.row_dimensions[r[0].row].height = 40
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 22
    ws.freeze_panes = "A4"

    # Risk summary sheet
    ws2 = wb.create_sheet("风险汇总")
    ws2.merge_cells("A1:D1")
    c = ws2["A1"]
    c.value = f"{platform_name} 主要合规风险汇总（按风险程度排序）"
    c.fill = HEADER
    c.font = Font(bold=True, color="FFFFFF", size=12)
    c.alignment = center
    c.border = border

    for i, h in enumerate(["优先级", "风险描述", "涉及条款位置", "整改建议"], 1):
        _sc(ws2, 2, i, h,
            fill=PatternFill("solid", fgColor="1F4E79"),
            font=Font(bold=True, color="FFFFFF", size=10),
            align=center)

    for i, row in enumerate(risks, 3):
        r_fill = RED if "高" in row[0] else (YELLOW if "中" in row[0] else GREEN)
        for j, val in enumerate(row, 1):
            _sc(ws2, i, j, val,
                fill=r_fill if j == 1 else None,
                font=Font(size=9, bold=(j == 1)),
                align=left)

    for i, w in enumerate([14, 46, 22, 46], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for r in ws2.iter_rows():
        ws2.row_dimensions[r[0].row].height = 52
    ws2.row_dimensions[1].height = 24
    ws2.row_dimensions[2].height = 20
    ws2.freeze_panes = "A3"

    wb.save(output_path)
    print(f"saved: {output_path}")


if __name__ == "__main__":
    # Smoke test
    sample_rows = [
        ["模型训练阶段", "收集处理与授权", "是否基于必需目的",
         "示例原文", "第X章", "存疑", "测试说明", "中"],
    ]
    sample_risks = [
        ["风险1（高）", "示例风险", "第X章", "示例建议"],
    ]
    generate_platform_report(
        "测试平台", "2026-01-01",
        "《个保法》《生成式AI服务管理暂行办法》",
        sample_rows, sample_risks,
        os.path.join(config.BASE, "test_report.xlsx")
    )
