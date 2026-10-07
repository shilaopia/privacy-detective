# -*- coding: utf-8 -*-
"""Stage 3: Generate 6-sheet Excel workbook from extracted brokerage contract clauses.

Reads combined_results.json (produced by Stage 2 extraction agents) and produces
a formatted .xlsx following make_report.py styling conventions.
"""

import json
import os
from datetime import datetime
from pathlib import Path
import config

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Styling constants (from make_report.py) ──────────────────────────────────
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
GRAY = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
SUBHDR_FILL = PatternFill("solid", fgColor="BDD7EE")
COLHDR_FILL = PatternFill("solid", fgColor="1F4E79")

thin = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADER_FONT = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
SUBHDR_FONT = Font(name="微软雅黑", size=10, color="2F5496")
COLHDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name="微软雅黑", size=9)
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="2F5496")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

CONFIDENCE_FILL = {"high": GREEN, "medium": YELLOW, "low": RED}
FOUND_FILL = {True: GREEN, False: RED}


def _sc(ws, row, col, value, fill=None, font=None, align=None):
    """Set cell value + style (reused from make_report.py pattern)."""
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    if font:
        c.font = font
    if align:
        c.alignment = align
    c.border = border
    return c


def write_title_block(ws, title, subtitle, col_count):
    """Write title row + subtitle row on a sheet (rows 1-2)."""
    _sc(ws, 1, 1, title, font=TITLE_FONT, align=LEFT)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    _sc(ws, 2, 1, subtitle, font=SUBHDR_FONT, align=LEFT)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 22


def write_header_row(ws, row, headers):
    """Write column headers with dark-blue fill, white bold font."""
    for ci, h in enumerate(headers, 1):
        _sc(ws, row, ci, h, fill=COLHDR_FILL, font=COLHDR_FONT, align=CENTER)
    ws.row_dimensions[row].height = 28


def write_data_row(ws, row, values, aligns=None):
    """Write one data row with border, font, alignment."""
    for ci, v in enumerate(values, 1):
        al = aligns[ci - 1] if aligns and ci - 1 < len(aligns) else LEFT
        _sc(ws, row, ci, v, font=DATA_FONT, align=al)
    ws.row_dimensions[row].height = 40


def auto_width(ws, min_w=8, max_w=50):
    """Auto-fit column widths (approximate for Chinese text)."""
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            if cell.value:
                # Chinese chars count as ~2
                text = str(cell.value)
                length = sum(2 if ord(c) > 127 else 1 for c in text)
                max_len = max(max_len, length)
        width = min(max(max_len * 1.1 + 2, min_w), max_w)
        ws.column_dimensions[col_letter].width = width


# ── Sheet builders ───────────────────────────────────────────────────────────

def build_sheet_overview(wb, data):
    """Sheet 1: 综合对比总览 — 7 contracts x 3 categories matrix."""
    ws = wb.create_sheet("综合对比总览")
    headers = [
        "合同名称", "合同期限", "期限类型", "续约机制", "通知期(天)",
        "基础分成比(甲:乙)", "结算周期", "附件引用", "备注"
    ]
    col_count = len(headers)

    write_title_block(ws, "演艺经纪合同 — 关键条款综合对比总览",
                      f"提取日期: {datetime.now().strftime('%Y-%m-%d')} | 共 {len(data)} 份合同",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    for ri, contract in enumerate(data):
        row = 4 + ri
        ct = contract.get("contract_term", {})
        rc = contract.get("renewal_conditions", {})
        ps = contract.get("profit_sharing", {})

        notes = []
        if ct.get("confidence") == "low":
            notes.append("期限需确认")
        if rc.get("confidence") == "low":
            notes.append("续约条款不完整")
        if ps.get("appendix_reference") and not ps.get("appendix_found_in_docx"):
            notes.append("附件未在docx中找到")

        values = [
            contract["contract_name"],
            f"{ct.get('start_date', '?')} 至 {ct.get('end_date', '?')} ({ct.get('term_years', '?')}年)" if ct.get("found") else "未约定",
            ct.get("term_type", "—"),
            rc.get("renewal_type", "—"),
            rc.get("notice_period_days", "—"),
            ps.get("base_split_ratio", "—"),
            ps.get("settlement_cycle", "—"),
            ps.get("appendix_reference", "无") if ps.get("appendix_reference") else "无",
            "; ".join(notes) if notes else "",
        ]
        aligns = [LEFT, LEFT, CENTER, CENTER, CENTER, CENTER, CENTER, LEFT, LEFT]
        write_data_row(ws, row, values, aligns)

        # Color-code confidence
        for cat, col_idx in [("contract_term", 2), ("renewal_conditions", 4), ("profit_sharing", 6)]:
            cat_data = contract.get(cat, {})
            if cat_data.get("found"):
                fill = CONFIDENCE_FILL.get(cat_data.get("confidence"), GRAY)
                ws.cell(row=row, column=col_idx).fill = fill

    auto_width(ws, max_w=40)
    return ws


def build_sheet_term(wb, data):
    """Sheet 2: 合同期限对比."""
    ws = wb.create_sheet("合同期限对比")
    headers = [
        "合同名称", "起算日期", "终止日期", "期限(年)", "期限类型",
        "提前终止条件", "原文摘录", "置信度"
    ]
    col_count = len(headers)

    write_title_block(ws, "合同期限条款对比",
                      "涵盖合作期限起止日、期限类型、提前终止条件",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    for ri, contract in enumerate(data):
        row = 4 + ri
        ct = contract.get("contract_term", {})
        if not ct.get("found"):
            values = [contract["contract_name"], "未约定", "未约定", "—", "—", "—", "—", "—"]
            write_data_row(ws, row, values)
            ws.cell(row=row, column=2).fill = RED
            continue

        values = [
            contract["contract_name"],
            ct.get("start_date", ""),
            ct.get("end_date", ""),
            ct.get("term_years", ""),
            ct.get("term_type", ""),
            ct.get("early_termination", ""),
            ct.get("source_text_raw", "")[:300],
            ct.get("confidence", ""),
        ]
        aligns = [LEFT, CENTER, CENTER, CENTER, CENTER, LEFT, LEFT_TOP, CENTER]
        write_data_row(ws, row, values, aligns)
        ws.cell(row=row, column=8).fill = CONFIDENCE_FILL.get(ct.get("confidence"), GRAY)

    auto_width(ws)
    # Widen source text column
    ws.column_dimensions[get_column_letter(7)].width = 50
    return ws


def build_sheet_renewal(wb, data):
    """Sheet 3: 续约条件对比."""
    ws = wb.create_sheet("续约条件对比")
    headers = [
        "合同名称", "续约方式", "通知期限(天)", "通知方式",
        "续约条件摘要", "续约后条款变更", "不续约/退出机制", "原文摘录", "置信度"
    ]
    col_count = len(headers)

    write_title_block(ws, "续约条件条款对比",
                      "涵盖续约机制、通知期限、条件及退出方式",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    for ri, contract in enumerate(data):
        row = 4 + ri
        rc = contract.get("renewal_conditions", {})
        if not rc.get("found"):
            values = [contract["contract_name"], "无约定", "—", "—", "—", "—", "—", "—", "—"]
            write_data_row(ws, row, values)
            ws.cell(row=row, column=2).fill = RED
            continue

        values = [
            contract["contract_name"],
            rc.get("renewal_type", ""),
            rc.get("notice_period_days", ""),
            rc.get("notice_method", ""),
            rc.get("renewal_conditions_summary", ""),
            "",  # renewal_term_change — may be nested in conditions
            rc.get("opt_out_mechanism", ""),
            rc.get("source_text_raw", "")[:300],
            rc.get("confidence", ""),
        ]
        aligns = [LEFT, CENTER, CENTER, CENTER, LEFT, CENTER, LEFT, LEFT_TOP, CENTER]
        write_data_row(ws, row, values, aligns)
        ws.cell(row=row, column=9).fill = CONFIDENCE_FILL.get(rc.get("confidence"), GRAY)

    auto_width(ws)
    ws.column_dimensions[get_column_letter(5)].width = 35
    ws.column_dimensions[get_column_letter(7)].width = 35
    ws.column_dimensions[get_column_letter(8)].width = 50
    return ws


def build_sheet_profit(wb, data):
    """Sheet 4: 利益分配对比 — the most important sheet with granular columns."""
    ws = wb.create_sheet("利益分配对比")
    headers = [
        "合同名称", "基础分成比(甲:乙)", "广告收入", "直播收入",
        "商务代言", "知识产权收益", "其他收入", "成本扣除方式",
        "结算周期", "结算截止日(天)", "税费承担", "保底/最低保障",
        "附件引用", "附件在docx", "原文摘录", "置信度"
    ]
    col_count = len(headers)

    write_title_block(ws, "利益分配与分成比例条款对比 ★重点★",
                      "涵盖基础分成、各收入类型差异化比例、结算规则、附件引用情况",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    for ri, contract in enumerate(data):
        row = 4 + ri
        ps = contract.get("profit_sharing", {})
        if not ps.get("found"):
            values = [contract["contract_name"]] + ["未约定"] * 14 + ["—"]
            write_data_row(ws, row, values)
            ws.cell(row=row, column=2).fill = RED
            continue

        # Build income-type-specific columns from split_adjustments
        adjustments = ps.get("split_adjustments", [])
        adj_map = {}
        for adj in adjustments:
            adj_map[adj.get("income_type", "")] = adj.get("ratio", "")

        # Check appendix
        appendix_in_docx = "是" if ps.get("appendix_found_in_docx") else "否"
        appendix_warn = ""
        if ps.get("appendix_reference") and not ps.get("appendix_found_in_docx"):
            appendix_warn = "⚠ 需人工核查"

        values = [
            contract["contract_name"],
            ps.get("base_split_ratio", ""),
            adj_map.get("广告收入", adj_map.get("广告", "")),
            adj_map.get("直播收入", adj_map.get("直播", adj_map.get("直播打赏", ""))),
            adj_map.get("商务代言", adj_map.get("商务", adj_map.get("代言", ""))),
            adj_map.get("知识产权收益", adj_map.get("知识产权", adj_map.get("IP", ""))),
            adj_map.get("其他收入", adj_map.get("其他", "")),
            ps.get("cost_deduction_policy", ""),
            ps.get("settlement_cycle", ""),
            ps.get("settlement_deadline_days", ""),
            ps.get("tax_handling", ""),
            ps.get("minimum_guarantee", ""),
            ps.get("appendix_reference", "无"),
            f"{appendix_in_docx} {appendix_warn}",
            ps.get("source_text_raw", "")[:300],
            ps.get("confidence", ""),
        ]
        aligns = [LEFT] + [CENTER] * 3 + [CENTER, CENTER, CENTER, LEFT, CENTER, CENTER, LEFT, LEFT, LEFT, CENTER, LEFT_TOP, CENTER]
        write_data_row(ws, row, values, aligns)
        ws.cell(row=row, column=16).fill = CONFIDENCE_FILL.get(ps.get("confidence"), GRAY)

        # Highlight appendix warnings
        if appendix_warn:
            ws.cell(row=row, column=14).fill = YELLOW

    auto_width(ws)
    ws.column_dimensions[get_column_letter(2)].width = 18
    ws.column_dimensions[get_column_letter(8)].width = 28
    ws.column_dimensions[get_column_letter(13)].width = 20
    ws.column_dimensions[get_column_letter(15)].width = 50
    return ws


def build_sheet_missing(wb, data):
    """Sheet 5: 缺失项汇总 — only records what is missing."""
    ws = wb.create_sheet("缺失项汇总")
    headers = ["合同名称", "缺失类别", "缺失项描述", "影响评估"]
    col_count = len(headers)

    write_title_block(ws, "条款缺失项汇总",
                      "以下列出各合同中未能提取的条款项，证明已做全覆盖审查",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    row = 4
    has_any = False
    for contract in data:
        name = contract["contract_name"]
        for cat_key, cat_label in [
            ("contract_term", "合同期限"),
            ("renewal_conditions", "续约条件"),
            ("profit_sharing", "利益分配"),
        ]:
            cat = contract.get(cat_key, {})
            if not cat.get("found"):
                has_any = True
                values = [
                    name, cat_label,
                    f"{cat_label}条款未在合同中找到",
                    "建议人工复查合同全文"
                ]
                write_data_row(ws, row, values)
                ws.cell(row=row, column=2).fill = RED
                row += 1
            # Also check for specific missing sub-fields
            if cat_key == "profit_sharing" and cat.get("found"):
                ps = cat
                if ps.get("appendix_reference") and not ps.get("appendix_found_in_docx"):
                    has_any = True
                    values = [
                        name, "利益分配—附件",
                        f"合同引用附件「{ps.get('appendix_reference')}」但docx内未找到",
                        "附件可能包含关键分成明细，强烈建议获取纸质版附件"
                    ]
                    write_data_row(ws, row, values)
                    ws.cell(row=row, column=2).fill = YELLOW
                    row += 1

    if not has_any:
        write_data_row(ws, row, ["（无缺失项）", "所有合同的全部三类条款均已成功提取", "", ""])
        ws.cell(row=row, column=1).fill = GREEN

    auto_width(ws)
    ws.column_dimensions[get_column_letter(3)].width = 45
    ws.column_dimensions[get_column_letter(4)].width = 45
    return ws


def build_sheet_audit(wb, data):
    """Sheet 6: 数据溯源 — audit trail."""
    ws = wb.create_sheet("数据溯源")
    headers = ["合同名称", "提取类别", "来源段落", "来源表格", "置信度", "提取时间", "警告/备注"]
    col_count = len(headers)

    write_title_block(ws, "数据溯源",
                      "记录每项提取的来源段落索引、置信度及警告信息",
                      col_count)
    write_header_row(ws, 3, headers)
    ws.freeze_panes = "A4"

    row = 4
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for contract in data:
        name = contract["contract_name"]
        for cat_key, cat_label in [
            ("contract_term", "合同期限"),
            ("renewal_conditions", "续约条件"),
            ("profit_sharing", "利益分配"),
        ]:
            cat = contract.get(cat_key, {})
            paragraphs = ", ".join(str(p) for p in cat.get("source_paragraphs", []))
            tables = ", ".join(str(t) for t in cat.get("source_tables", []))
            warnings = ""
            if cat_key == "profit_sharing" and cat.get("appendix_reference") and not cat.get("appendix_found_in_docx"):
                warnings = f"附件未在docx中找到: {cat.get('appendix_reference')}"

            values = [
                name, cat_label,
                paragraphs if paragraphs else "未找到",
                tables if tables else "—",
                cat.get("confidence", "—"),
                now,
                warnings,
            ]
            write_data_row(ws, row, values)
            ws.cell(row=row, column=5).fill = CONFIDENCE_FILL.get(cat.get("confidence"), GRAY)
            row += 1

    auto_width(ws)
    ws.column_dimensions[get_column_letter(3)].width = 30
    ws.column_dimensions[get_column_letter(4)].width = 20
    ws.column_dimensions[get_column_letter(7)].width = 40
    return ws


# ── Main ─────────────────────────────────────────────────────────────────────

def generate_report(input_json: str, output_path: str):
    """Generate the 6-sheet Excel workbook from extracted JSON data."""
    with open(input_json, "r", encoding="utf-8") as f:
        data = json.load(f)

    if isinstance(data, dict):
        # Could be {contracts: [...]} or directly a list
        contracts = data.get("contracts", data.get("data", data))
        if isinstance(contracts, dict):
            contracts = list(contracts.values())
    else:
        contracts = data

    if not isinstance(contracts, list):
        raise ValueError(f"Unexpected data format: expected list, got {type(contracts)}")

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    build_sheet_overview(wb, contracts)
    build_sheet_term(wb, contracts)
    build_sheet_renewal(wb, contracts)
    build_sheet_profit(wb, contracts)
    build_sheet_missing(wb, contracts)
    build_sheet_audit(wb, contracts)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Report saved to: {output_path}")
    print(f"  Sheets: {wb.sheetnames}")
    print(f"  Contracts: {len(contracts)}")


if __name__ == "__main__":
    import sys

    input_json = sys.argv[1] if len(sys.argv) > 1 else \
        os.path.join(config.BASE, "_brokerage_extract", "_combined_results.json")
    output_path = sys.argv[2] if len(sys.argv) > 2 else \
        os.path.join(config.BASE, "brokerage_excels", "经纪约关键条款提取对比_20260710.xlsx")

    generate_report(input_json, output_path)
