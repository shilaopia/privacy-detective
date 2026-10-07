# -*- coding: utf-8 -*-
"""第三轮产物生成:
A. privacy_excels\\制度标尺v7-细则级验证-3平台.xlsx
   - sheet「候选细则明细」: 149 候选细则 × 3 平台(评估+说明)
   - sheet「条款普遍达到度」: 58 条款 × 25 目标合规分布 + 候选标记
B. privacy_excels\\制度标尺v7.1-精简版.xlsx
   - sheet「全量细则标记」: 449 细则逐条 保留/降权复核/移除 + 依据
   - sheet「摘要」: 处理结果计数与维度分布
另附 merge_parts(key): 合并 review_details/parts/{key}__p*.json → review_details/{key}.json
"""
import json
import os
import sys
import glob
import config
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = config.BASE
TARGETS = ["PackyCode", "LiteLLM", "Helicone"]
OUT_A = os.path.join(BASE, "privacy_excels", "制度标尺v7-细则级验证-3平台.xlsx")
OUT_B = os.path.join(BASE, "privacy_excels", "制度标尺v7.1-精简版.xlsx")

FILL = {
    "合规": PatternFill("solid", fgColor="C6EFCE"),
    "存疑": PatternFill("solid", fgColor="FFEB9C"),
    "不合规": PatternFill("solid", fgColor="FFC7CE"),
    "未提及": PatternFill("solid", fgColor="D9D9D9"),
}
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def merge_parts(key):
    parts_dir = os.path.join(BASE, "review_details", "parts")
    merged = {}
    for p in sorted(glob.glob(os.path.join(parts_dir, f"{key}__p*.json"))):
        merged.update(json.load(open(p, encoding="utf-8")))
    focus = json.load(open(os.path.join(BASE, "focus_details.json"), encoding="utf-8"))
    lack = [str(d["id"]) for d in focus if str(d["id"]) not in merged]
    if lack:
        raise SystemExit(f"{key} 缺细则: {lack}")
    out = os.path.join(BASE, "review_details", key + ".json")
    json.dump(merged, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"已合并 {key}: {len(merged)} 条")
    return merged


def _verdict_cell(ws, r, c, v):
    cell = ws.cell(r, c, v)
    cell.fill = FILL.get(v, FILL["未提及"])
    cell.alignment = CENTER
    cell.border = THIN


def main():
    focus = json.load(open(os.path.join(BASE, "focus_details.json"), encoding="utf-8"))
    alld = json.load(open(os.path.join(BASE, "baseline_v7_details.json"), encoding="utf-8"))
    prev = json.load(open(os.path.join(BASE, "clause_prevalence.json"), encoding="utf-8"))
    reviews = {}
    for key in TARGETS:
        p = os.path.join(BASE, "review_details", key + ".json")
        if not os.path.exists(p):
            raise SystemExit(f"缺少: {p} (先运行 python gen_detail_review.py merge)")
        reviews[key] = json.load(open(p, encoding="utf-8"))
    focus_ids = {d["id"] for d in focus}

    # ---------- A. 验证分析表 ----------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "候选细则明细"
    head = ["细则ID", "一级维度", "核心条款", "义务细则", "出处"]
    ws.append(head + [t for t in TARGETS for t in (t, "")])
    ws.append(["", "", "", "", ""] + ["评估", "说明"] * 3)
    for c in range(1, 6):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for i in range(3):
        ws.merge_cells(start_row=1, start_column=6 + i * 2, end_row=1, end_column=7 + i * 2)
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=11):
        for cell in row:
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = CENTER
            cell.border = THIN
    r = 3
    for d in focus:
        ws.cell(r, 1, d["id"]).alignment = CENTER
        ws.cell(r, 2, d["dim1"])
        ws.cell(r, 3, d["clause"]).alignment = WRAP
        ws.cell(r, 4, d["text"]).alignment = WRAP
        ws.cell(r, 5, d["source"]).alignment = WRAP
        for i, key in enumerate(TARGETS):
            rd = reviews[key].get(str(d["id"]), {})
            _verdict_cell(ws, r, 6 + i * 2, rd.get("评估", "未提及"))
            nc = ws.cell(r, 7 + i * 2, rd.get("说明", ""))
            nc.alignment = WRAP
            nc.border = THIN
        for c in range(1, 6):
            ws.cell(r, c).border = THIN
        r += 1
    ws.freeze_panes = "F3"
    for col, w in zip("ABCDEFGHIJK", (7, 14, 20, 40, 22, 6, 36, 6, 36, 6, 36)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("条款普遍达到度")
    ws2.append(["核心条款", "合规", "存疑", "不合规", "未提及", "25目标合规率", "入选候选集"])
    for cell in ws2[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
    focus_clauses = {d["clause"] for d in focus}
    for cl, s in sorted(prev.items(), key=lambda kv: -kv[1]["合规"]):
        ws2.append([cl, s["合规"], s["存疑"], s["不合规"], s["未提及"],
                    f"{s['合规']/25:.0%}", "是" if cl in focus_clauses else ""])
    ws2.column_dimensions["A"].width = 40
    for col in "BCDEFG":
        ws2.column_dimensions[col].width = 12
    wb.save(OUT_A)
    print("已生成:", OUT_A)

    # ---------- B. 精简版标尺 ----------
    wb2 = openpyxl.Workbook()
    ws3 = wb2.active
    ws3.title = "全量细则标记"
    ws3.append(["细则ID", "一级维度", "核心条款", "义务细则", "出处", "GAI特有",
                "条款25目标合规数", "三家验证", "处理", "依据"])
    for cell in ws3[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN

    FILL_KEEP = PatternFill("solid", fgColor="C6EFCE")
    FILL_DROP = PatternFill("solid", fgColor="D9D9D9")
    FILL_HOLD = PatternFill("solid", fgColor="FFEB9C")
    cnt = Counter()
    r = 2
    for d in alld:
        in_cand = d["id"] in focus_ids
        vs = [reviews[k].get(str(d["id"]), {}).get("评估", "未提及") for k in TARGETS]
        if not in_cand:
            action, why, fill = "保留", "区分度高(条款合规数<15/25)或GAI特有", FILL_KEEP
        elif all(v == "合规" for v in vs):
            action, why, fill = "移除", "普遍达到且尾部3平台全部合规,验证通过", FILL_DROP
        elif any(v in ("不合规", "未提及") for v in vs):
            action, why, fill = "保留", f"尾部平台失守({'/'.join(vs)}),有区分价值", FILL_KEEP
        else:
            action, why, fill = "降权复核", f"尾部平台判定存疑({'/'.join(vs)})", FILL_HOLD
        cnt[action] += 1
        row = [d["id"], d["dim1"], d["clause"], d["text"], d["source"],
               "是" if d["gai"] else "", prev.get(d["clause"], {}).get("合规", ""),
               "/".join(vs) if in_cand else "", action, why]
        ws3.append(row)
        ac = ws3.cell(r, 9)
        ac.fill = fill
        ac.alignment = CENTER
        for c in range(1, 11):
            ws3.cell(r, c).border = THIN
            if c in (3, 4, 5, 10):
                ws3.cell(r, c).alignment = WRAP
        r += 1
    ws3.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJ", (7, 14, 20, 44, 20, 8, 10, 16, 10, 34)):
        ws3.column_dimensions[col].width = w

    ws4 = wb2.create_sheet("摘要")
    ws4.append(["处理", "细则数"])
    for cell in ws4[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    for a in ("保留", "降权复核", "移除"):
        ws4.append([a, cnt[a]])
    ws4.append(["合计", sum(cnt.values())])
    ws4.append([])
    ws4.append(["验证逻辑", "候选移除集=普遍达到(条款≥15/25合规)∩非GAI; 候选集内尾部3平台全合规→移除; 有存疑→降权复核; 有未提及/不合规→保留"])
    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 90
    wb2.save(OUT_B)
    print("已生成:", OUT_B, dict(cnt))


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "merge":
        for key in TARGETS:
            merge_parts(key)
    else:
        main()
