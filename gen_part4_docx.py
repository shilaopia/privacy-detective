# -*- coding: utf-8 -*-
"""将结项书第四部分正文草稿（markdown）渲染为独立 docx。
正文宋体小四、首行缩进2字符、1.5倍行距；标题黑体；图表居中配题注。
"""
import re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(BASE, "结项书第四部分_正文草稿.md")
PIC = os.path.join(BASE, "pictures", "大创合规图")
OUT = os.path.join(BASE, "结项书第四部分_平台规则文本合法合规研究.docx")

FIGS = {
    "图4-1": os.path.join(PIC, "微信图片_20260919211751_1032_12.png"),
    "图4-2": os.path.join(PIC, "微信图片_2026-09-26_094745_163.png"),
    "图4-3": os.path.join(PIC, "修订版", "图4-3_三层判定结构堆叠图.png"),
    "图4-4": os.path.join(PIC, "微信图片_2026-09-26_094757_346.png"),
    "图4-5": os.path.join(PIC, "微信图片_2026-09-26_094802_317.png"),
    "图4-6": os.path.join(PIC, "微信图片_2026-09-26_094750_979.png"),
    "图4-7": os.path.join(PIC, "修订版", "图4-7_七维度合规率均值图.png"),
    "图4-8": os.path.join(PIC, "微信图片_20260919203249_1031_12.png"),
    "图4-9": os.path.join(PIC, "修订版", "图4-9_跨境4-3系列分层合规率.png"),
    "附表": os.path.join(PIC, "微信图片_2026-09-26_094823_860.png"),
}

def set_font(run, ascii_f="Times New Roman", ea="宋体", size=12, bold=False):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)

def body_para(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        pf.first_line_indent = Pt(24)  # 2×小四
    pf.space_after = Pt(0)
    # 处理 **加粗** 段
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = p.add_run(seg)
        set_font(r, size=12, bold=(i % 2 == 1))
    return p

def head_para(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(13 if level == 0 else 6)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    sizes = {0: 16, 1: 14, 2: 12}
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, ea="黑体", size=sizes[level], bold=True)
    return p

def caption(doc, text, before=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if before else 3)
    p.paragraph_format.space_after = Pt(3 if before else 6)
    r = p.add_run(text)
    set_font(r, ea="宋体", size=10.5, bold=True)
    return p

def add_fig(doc, key, cap):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width = Cm(15.5) if key == "图4-7" else Cm(14)
    p.add_run().add_picture(FIGS[key], width=width)
    caption(doc, cap, before=False)

def add_table(doc, cap, rows):
    caption(doc, cap, before=True)
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci > 0) else WD_ALIGN_PARAGRAPH.LEFT
            r = para.add_run(val.strip())
            set_font(r, size=10.5, bold=(ri == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc = Document()
# 页面：A4，常规页边距
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.8)

lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i].strip()
    if not ln or ln == "---":
        i += 1
        continue
    if ln.startswith("### "):
        head_para(doc, ln[4:], 2)
    elif ln.startswith("## "):
        head_para(doc, ln[3:], 1)
    elif ln.startswith("# "):
        head_para(doc, ln[2:], 0)
    elif ln.startswith("【图") or ln.startswith("【附表"):
        m = re.match(r"【(.+?)】", ln)
        cap = m.group(1)
        key = cap.split()[0]
        if "此处插入" in ln or key == "附表":
            add_fig(doc, key, cap)
        else:
            caption(doc, cap)
    elif ln.startswith("【表"):
        cap = re.match(r"【(.+?)】", ln).group(1)
        rows = []
        i += 1
        while i < len(lines) and not lines[i].strip():
            i += 1
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c.strip()) <= set("-: ") for c in cells):
                rows.append(cells)
            i += 1
        add_table(doc, cap, rows)
        continue
    else:
        body_para(doc, ln)
    i += 1

doc.save(OUT)
print("OK", OUT)
