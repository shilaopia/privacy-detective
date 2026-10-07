# -*- coding: utf-8 -*-
"""第二轮:把 R2_* 审查结果并入现有 制度标尺v7-横向对比.xlsx。

新增 sheet「API中转与聚合」(58 条款 × 14 目标 × 2 列) 与「汇总-API」,
样式与第一轮主表一致;不改动第一轮已有 sheet。
另附 merge_parts():把 review_json/parts/R2_{key}__p{1..7}.json 合并为 review_json/R2_{key}.json。
"""
import json
import os
import sys
import glob
import openpyxl
import config
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = config.BASE
OUT = os.path.join(BASE, "privacy_excels", "制度标尺v7-横向对比.xlsx")

# 目标列序(与分类表一致:9 聚合 + 2 框架 + 3 运营商)
TARGETS = [
    ("R2_OpenRouter", "OpenRouter"), ("R2_Portkey", "Portkey"),
    ("R2_CloudflareAG", "Cloudflare\nAI Gateway"), ("R2_VercelAG", "Vercel\nAI Gateway"),
    ("R2_Helicone", "Helicone"), ("R2_LiteLLM", "LiteLLM"),
    ("R2_PackyCode", "PackyCode"), ("R2_API易", "API易"), ("R2_UniAPI", "UniAPI"),
    ("R2_OneAPI", "OneAPI\n(框架)"), ("R2_NewAPI", "NewAPI\n(框架)"),
    ("R2_移动MoMA", "移动\nMoMA"), ("R2_电信天翼云", "电信\n天翼云"), ("R2_联通星罗", "联通\n星罗"),
]

FILL = {
    "合规": PatternFill("solid", fgColor="C6EFCE"),
    "存疑": PatternFill("solid", fgColor="FFEB9C"),
    "不合规": PatternFill("solid", fgColor="FFC7CE"),
    "未提及": PatternFill("solid", fgColor="D9D9D9"),
}
DIM_FILL = PatternFill("solid", fgColor="DEEAF1")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def merge_parts(key):
    """合并分片 -> review_json/{key}.json,返回缺失分片列表(空=完整)。"""
    parts_dir = os.path.join(BASE, "review_json", "parts")
    baseline = json.load(open(os.path.join(BASE, "baseline_v7.json"), encoding="utf-8"))
    missing = []
    merged = {}
    for n in range(1, 8):
        p = os.path.join(parts_dir, f"{key}__p{n}.json")
        if not os.path.exists(p):
            missing.append(n)
            continue
        merged.update(json.load(open(p, encoding="utf-8")))
    if missing:
        return missing
    lack = [c["clause"] for c in baseline if c["clause"] not in merged]
    extra = [k for k in merged if k not in {c["clause"] for c in baseline}]
    if lack:
        raise SystemExit(f"{key} 合并后缺条款: {lack}; 多余键: {extra}")
    out = os.path.join(BASE, "review_json", key + ".json")
    json.dump(merged, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"已合并 {key}: {len(merged)} 条 -> {out}")
    return []


def main():
    baseline = json.load(open(os.path.join(BASE, "baseline_v7.json"), encoding="utf-8"))
    reviews = {}
    for key, _ in TARGETS:
        p = os.path.join(BASE, "review_json", key + ".json")
        if not os.path.exists(p):
            raise SystemExit(f"缺少审查结果: {p}")
        reviews[key] = json.load(open(p, encoding="utf-8"))

    wb = openpyxl.load_workbook(OUT)
    for name in ("API中转与聚合", "汇总-API"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("API中转与聚合")
    fixed = ["一级维度", "二级维度", "核心对标条款", "义务\n细则数"]
    ws.append(fixed + [t for _, name in TARGETS for t in (name, "")])
    ws.append(["", "", "", ""] + ["评估", "合规说明"] * len(TARGETS))
    for c in range(1, 5):
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(1, c).font = HEAD_FONT
        ws.cell(1, c).alignment = CENTER
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for i, (_, name) in enumerate(TARGETS):
        c1 = 5 + i * 2
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c1 + 1)
        for r in (1, 2):
            ws.cell(r, c1).fill = HEAD_FILL
            ws.cell(r, c1 + 1).fill = HEAD_FILL
        ws.cell(1, c1).font = HEAD_FONT
        ws.cell(1, c1).alignment = CENTER
        ws.cell(2, c1).font = HEAD_FONT
        ws.cell(2, c1 + 1).font = HEAD_FONT
        ws.cell(2, c1).alignment = CENTER
        ws.cell(2, c1 + 1).alignment = CENTER
    for c in range(1, 5):
        ws.cell(2, c).fill = HEAD_FILL

    r = 3
    dim1_rows = {}
    for cl in baseline:
        dim1_rows.setdefault(cl["dim1"], []).append(r)
        ws.cell(r, 1, cl["dim1"])
        ws.cell(r, 2, cl["dim2"] or "—")
        ws.cell(r, 3, cl["clause"])
        ws.cell(r, 4, cl["detail_count"])
        for i, (key, _) in enumerate(TARGETS):
            cell_data = reviews[key].get(cl["clause"], {})
            verdict = cell_data.get("评估", "未提及")
            note = cell_data.get("合规说明", "")
            cv = ws.cell(r, 5 + i * 2, verdict)
            cv.fill = FILL.get(verdict, FILL["未提及"])
            cv.alignment = CENTER
            ws.cell(r, 6 + i * 2, note).alignment = WRAP
        r += 1

    for dim, rows in dim1_rows.items():
        if len(rows) > 1:
            ws.merge_cells(start_row=rows[0], start_column=1, end_row=rows[-1], end_column=1)
        for rr in rows:
            ws.cell(rr, 1).fill = DIM_FILL
            ws.cell(rr, 1).alignment = CENTER

    last_col = 4 + 2 * len(TARGETS)
    for row in ws.iter_rows(min_row=1, max_row=r - 1, max_col=last_col):
        for cell in row:
            cell.border = THIN
            if 2 <= cell.column <= 3 and cell.row >= 3:
                cell.alignment = WRAP
            elif cell.column == 4 and cell.row >= 3:
                cell.alignment = CENTER

    ws.freeze_panes = "E3"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 7
    for i in range(len(TARGETS)):
        ws.column_dimensions[get_column_letter(5 + i * 2)].width = 6
        ws.column_dimensions[get_column_letter(6 + i * 2)].width = 42
    ws.row_dimensions[1].height = 30
    for rr in range(3, r):
        ws.row_dimensions[rr].height = 60

    ws2 = wb.create_sheet("汇总-API")
    ws2.append(["目标", "合规", "存疑", "不合规", "未提及", "合规率(合规/58)"])
    for cell in ws2[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
    for key, name in TARGETS:
        counts = {v: 0 for v in FILL}
        for cl in baseline:
            v = reviews[key].get(cl["clause"], {}).get("评估", "未提及")
            counts[v if v in counts else "未提及"] += 1
        ws2.append([name.replace("\n", ""), counts["合规"], counts["存疑"],
                    counts["不合规"], counts["未提及"], f"{counts['合规']/len(baseline):.0%}"])
    ws2.column_dimensions["A"].width = 16
    for col in "BCDEF":
        ws2.column_dimensions[col].width = 12

    wb.save(OUT)
    print("已更新:", OUT)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "merge":
        for key, _ in TARGETS:
            miss = merge_parts(key)
            if miss:
                print(f"{key} 缺分片: {miss}")
    else:
        main()
