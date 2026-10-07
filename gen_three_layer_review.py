# -*- coding: utf-8 -*-
"""三层样本 v7 制度标尺核验 — 最终合成 Excel（64 目标版）。
读 review_json/{key}.json(58条款) × 64 目标,产出
privacy_excels/制度标尺v7-三层样本核验.xlsx,含:
  核验主表(58条款×64目标,层→子类两级表头)
  数据汇总(层/子类/目标/端别/四项计数/四项比率/主要不合规条款 + 子类与层小计)
  条款普遍达到度(58条款 总体及分层比率)
  重点分析(火山方舟/OpenRouter,读 review_json/focus_{key}.json,无则占位)
  模型层训练专项(读 review_json/training_focus.json,无则占位)
运行: python gen_three_layer_review.py
"""
import json
import os
import config
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = config.BASE
OUT = os.path.join(BASE, "privacy_excels", "制度标尺v7-三层样本核验.xlsx")

# (层, [(子类, [(key, 显示名, 端别), ...]), ...]) — key 对应 review_json/{key}.json
LAYERS = [
    ("应用层", [
        ("AI智能体", [
            ("扣子Coze_ToC", "扣子Coze", "To-C"), ("腾讯元器_ToC", "腾讯元器", "To-C"),
            ("AutoGLM_ToC", "AutoGLM", "To-C"), ("支小宝_ToC", "支小宝", "To-C"),
        ]),
        ("传统内嵌AI", [
            ("京东AI购_ToC", "京东AI购", "To-C"), ("淘宝_ToC", "淘宝", "To-C"),
            ("抖音_ToC", "抖音", "To-C"), ("小红书_ToC", "小红书", "To-C"),
            ("小美App_ToC", "小美App", "To-C"), ("滴滴送货AI助手_ToC", "滴滴送货AI助手", "To-C"),
            ("高德地图_ToC", "高德地图", "To-C"), ("支付宝_ToC", "支付宝", "To-C"),
            ("夸克_ToC", "夸克", "To-C"), ("超级小爱_ToC", "超级小爱", "To-C"),
            ("Keep_ToC", "Keep", "To-C"), ("京东健康_ToC", "京东健康", "To-C"),
            ("飞书智能伙伴_ToC", "飞书智能伙伴", "To-C"), ("腾讯文档_ToC", "腾讯文档", "To-C"),
            ("钉钉_ToC", "钉钉", "To-C"),
        ]),
    ]),
    ("模型层", [
        ("大模型厂商", [
            ("DeepSeek_ToC", "DeepSeek", "To-C"), ("DeepSeek_ToB", "DeepSeek", "To-B"),
            ("豆包_ToC", "豆包", "To-C"), ("豆包_ToB", "豆包", "To-B"),
            ("Kimi_ToC", "Kimi", "To-C"), ("Kimi_ToB", "Kimi", "To-B"),
            ("千问_ToC", "千问", "To-C"), ("千问_ToB", "千问", "To-B"),
            ("元宝_ToC", "元宝", "To-C"), ("元宝_ToB", "元宝", "To-B"),
            ("智谱清言_ToC", "智谱清言", "To-C"), ("智谱清言_ToB", "智谱清言", "To-B"),
            ("文心一言_ToC", "文心一言", "To-C"), ("文心一言_ToB", "文心一言", "To-B"),
            ("混元_ToC", "混元", "To-C"), ("混元_ToB", "混元", "To-B"),
            ("minimax_ToC", "minimax", "To-C"), ("minimax_ToB", "minimax", "To-B"),
            ("讯飞星火_ToC", "讯飞星火", "To-C"), ("讯飞星火_ToB", "讯飞星火", "To-B"),
            ("零一万物_ToC", "零一万物", "To-C"), ("零一万物_ToB", "零一万物", "To-B"),
            ("商汤_ToC", "商汤", "To-C"), ("商汤_ToB", "商汤", "To-B"),
        ]),
    ]),
    ("分发层", [
        ("境内-运营商", [
            ("移动MoMA", "中国移动MoMA", "To-B"), ("电信天翼云", "中国电信天翼云", "To-B"),
            ("联通星罗", "中国联通星罗", "To-B"),
        ]),
        ("境内-云厂商MaaS", [
            ("火山方舟", "火山方舟★", "To-B"), ("阿里云百炼", "阿里云百炼", "To-B"),
            ("百度千帆", "百度智能云千帆", "To-B"), ("腾讯云LKE", "腾讯云知识引擎", "To-B"),
            ("华为云ModelArts", "华为云ModelArts", "To-B"),
        ]),
        ("境内-API中转聚合", [
            ("PackyCode", "PackyAPI/PackyCode", "To-B"), ("API易", "API易", "To-B"),
            ("UniAPI", "UniAPI", "To-B"), ("OneAPI", "OneAPI(框架)", "To-B"),
            ("NewAPI", "NewAPI(框架)", "To-B"),
        ]),
        ("境外-代表平台", [
            ("OpenRouter", "OpenRouter★", "To-B"), ("TogetherAI", "Together AI", "To-B"),
            ("FireworksAI", "Fireworks AI", "To-B"),
        ]),
        ("境外-Gateway", [
            ("Portkey", "Portkey AI", "To-B"), ("CloudflareAG", "Cloudflare AI Gateway", "To-B"),
            ("VercelAG", "Vercel AI Gateway", "To-B"), ("Helicone", "Helicone", "To-B"),
            ("LiteLLM", "LiteLLM(开源)", "To-B"),
        ]),
    ]),
]

FOCUS_TARGETS = [("火山方舟", "火山方舟（字节跳动 MaaS）重点分析"),
                 ("OpenRouter", "OpenRouter（境外模型分发代表）重点分析")]

FILL = {
    "合规": PatternFill("solid", fgColor="C6EFCE"),
    "存疑": PatternFill("solid", fgColor="FFEB9C"),
    "不合规": PatternFill("solid", fgColor="FFC7CE"),
    "未提及": PatternFill("solid", fgColor="D9D9D9"),
}
DIM_FILL = PatternFill("solid", fgColor="DEEAF1")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
SUB_FILL = PatternFill("solid", fgColor="8EAADB")
LAYER_FILLS = {
    "应用层": PatternFill("solid", fgColor="2E75B6"),
    "模型层": PatternFill("solid", fgColor="548235"),
    "分发层": PatternFill("solid", fgColor="BF8F00"),
}
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def flat_targets():
    return [(layer, sub, key, name, ep)
            for layer, subs in LAYERS for sub, ts in subs for key, name, ep in ts]


def load_all():
    baseline = json.load(open(os.path.join(BASE, "baseline_v7.json"), encoding="utf-8"))
    reviews, missing = {}, []
    for _, _, key, _, _ in flat_targets():
        p = os.path.join(BASE, "review_json", key + ".json")
        if os.path.exists(p):
            reviews[key] = json.load(open(p, encoding="utf-8"))
        else:
            missing.append(key)
    return baseline, reviews, missing


def sheet_main(wb, baseline, reviews):
    ws = wb.create_sheet("核验主表")
    flat = flat_targets()
    n2 = len(flat) * 2
    # 四行表头: 层 / 子类 / 目标 / 评估+合规说明
    ws.append(["一级维度", "二级维度", "核心对标条款", "义务\n细则数"] + [""] * n2)
    ws.append([""] * 4 + [""] * n2)
    ws.append([""] * 4 + [""] * n2)
    ws.append([""] * 4 + ["评估", "合规说明"] * len(flat))
    for c in range(1, 5):
        ws.merge_cells(start_row=1, start_column=c, end_row=4, end_column=c)
        for r in range(1, 5):
            ws.cell(r, c).fill = HEAD_FILL
            ws.cell(r, c).font = HEAD_FONT
        ws.cell(1, c).alignment = CENTER
    col = 5
    for layer, subs in LAYERS:
        layer_span = sum(len(ts) for _, ts in subs) * 2
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + layer_span - 1)
        cell = ws.cell(1, col, f"{layer}（{layer_span // 2}）")
        cell.fill = LAYER_FILLS[layer]
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        for cc in range(col, col + layer_span):
            ws.cell(1, cc).fill = LAYER_FILLS[layer]
        for sub, ts in subs:
            span = len(ts) * 2
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + span - 1)
            sc = ws.cell(2, col, sub)
            sc.fill = SUB_FILL
            sc.font = Font(bold=True)
            sc.alignment = CENTER
            for cc in range(col, col + span):
                ws.cell(2, cc).fill = SUB_FILL
            for i, (key, name, ep) in enumerate(ts):
                c1 = col + i * 2
                ws.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c1 + 1)
                nm = ws.cell(3, c1, f"{name}\n{ep}")
                nm.fill = HEAD_FILL
                nm.font = HEAD_FONT
                nm.alignment = CENTER
                ws.cell(3, c1 + 1).fill = HEAD_FILL
                for cc in (c1, c1 + 1):
                    ws.cell(4, cc).fill = HEAD_FILL
                    ws.cell(4, cc).font = HEAD_FONT
                    ws.cell(4, cc).alignment = CENTER
            col += span

    r = 5
    dim1_rows = {}
    for cl in baseline:
        dim1_rows.setdefault(cl["dim1"], []).append(r)
        ws.cell(r, 1, cl["dim1"])
        ws.cell(r, 2, cl["dim2"] or "—")
        ws.cell(r, 3, cl["clause"])
        ws.cell(r, 4, cl["detail_count"])
        for i, (_, _, key, _, _) in enumerate(flat):
            cell_data = reviews[key].get(cl["clause"], {})
            verdict = cell_data.get("评估", "未提及")
            cv = ws.cell(r, 5 + i * 2, verdict)
            cv.fill = FILL.get(verdict, FILL["未提及"])
            cv.alignment = CENTER
            ws.cell(r, 6 + i * 2, cell_data.get("合规说明", "")).alignment = WRAP
        r += 1

    for dim, rows in dim1_rows.items():
        if len(rows) > 1:
            ws.merge_cells(start_row=rows[0], start_column=1, end_row=rows[-1], end_column=1)
        for rr in rows:
            ws.cell(rr, 1).fill = DIM_FILL
            ws.cell(rr, 1).alignment = CENTER

    last_col = 4 + 2 * len(flat)
    for row in ws.iter_rows(min_row=1, max_row=r - 1, max_col=last_col):
        for cell in row:
            cell.border = THIN
            if 2 <= cell.column <= 3 and cell.row >= 5:
                cell.alignment = WRAP
            elif cell.column == 4 and cell.row >= 5:
                cell.alignment = CENTER

    ws.freeze_panes = "E5"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 7
    for i in range(len(flat)):
        ws.column_dimensions[get_column_letter(5 + i * 2)].width = 6
        ws.column_dimensions[get_column_letter(6 + i * 2)].width = 42
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 30
    for rr in range(5, r):
        ws.row_dimensions[rr].height = 60


def counts_of(review, baseline):
    counts = {v: 0 for v in FILL}
    bad = []
    for cl in baseline:
        v = review.get(cl["clause"], {}).get("评估", "未提及")
        counts[v if v in counts else "未提及"] += 1
        if v == "不合规":
            bad.append(cl["clause"].split("-", 3)[-1])
    return counts, bad


def summary_row(ws, label_cells, counts, denom_n, n_targets, bad, head=False):
    ws.append(label_cells + [counts["合规"], counts["存疑"], counts["不合规"], counts["未提及"],
                             counts["合规"] / denom_n, counts["存疑"] / denom_n,
                             counts["不合规"] / denom_n, counts["未提及"] / denom_n, bad])
    r = ws.max_row
    for c in range(9, 13):
        ws.cell(r, c).number_format = "0.0%"
    for c in range(1, 14):
        ws.cell(r, c).border = THIN
        ws.cell(r, c).alignment = WRAP if c in (2, 14) else CENTER
        if head:
            ws.cell(r, c).fill = DIM_FILL
            ws.cell(r, c).font = Font(bold=True)


def sheet_summary(wb, baseline, reviews):
    ws = wb.create_sheet("数据汇总")
    head = ["层", "子类", "目标", "端别", "合规", "存疑", "不合规", "未提及",
            "合规率", "存疑率", "不合规率", "未提及率", "主要不合规条款"]
    ws.append(head)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN
    n = len(baseline)
    for layer, subs in LAYERS:
        layer_tot = {v: 0 for v in FILL}
        layer_cnt = 0
        for sub, ts in subs:
            sub_tot = {v: 0 for v in FILL}
            for key, name, ep in ts:
                counts, bad = counts_of(reviews[key], baseline)
                summary_row(ws, [layer, sub, name, ep], counts, n, 1, "；".join(bad))
                for v in sub_tot:
                    sub_tot[v] += counts[v]
            summary_row(ws, [layer, f"{sub}小计", f"（{len(ts)}目标）", ""],
                        sub_tot, n * len(ts), len(ts), "", head=True)
            for v in layer_tot:
                layer_tot[v] += sub_tot[v]
            layer_cnt += len(ts)
        summary_row(ws, [f"{layer}合计", "", f"（{layer_cnt}目标）", ""],
                    layer_tot, n * layer_cnt, layer_cnt, "", head=True)
    widths = {"A": 12, "B": 16, "C": 20, "D": 8, "E": 7, "F": 7, "G": 8, "H": 8,
              "I": 9, "J": 9, "K": 9, "L": 9, "M": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def sheet_prevalence(wb, baseline, reviews):
    ws = wb.create_sheet("条款普遍达到度")
    head = ["一级维度", "二级维度", "核心对标条款",
            "总体合规率", "总体存疑率", "总体不合规率", "总体未提及率",
            "应用层合规率", "模型层合规率", "分发层合规率"]
    ws.append(head)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN
    layer_keys = {layer: [k for _, ts in subs for k, _, _ in ts] for layer, subs in LAYERS}
    all_keys = [k for ks in layer_keys.values() for k in ks]
    for cl in baseline:
        def rate(keys, val):
            return sum(1 for k in keys
                       if reviews[k].get(cl["clause"], {}).get("评估", "未提及") == val) / len(keys)
        ws.append([cl["dim1"], cl["dim2"] or "—", cl["clause"],
                   rate(all_keys, "合规"), rate(all_keys, "存疑"),
                   rate(all_keys, "不合规"), rate(all_keys, "未提及"),
                   rate(layer_keys["应用层"], "合规"),
                   rate(layer_keys["模型层"], "合规"),
                   rate(layer_keys["分发层"], "合规")])
        r = ws.max_row
        for c in range(4, 11):
            ws.cell(r, c).number_format = "0.0%"
            ws.cell(r, c).alignment = CENTER
            ws.cell(r, c).border = THIN
        for c in range(1, 4):
            ws.cell(r, c).alignment = WRAP
            ws.cell(r, c).border = THIN
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 34
    for c in range(4, 11):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.freeze_panes = "D2"


def sheet_focus(wb):
    """重点分析: 火山方舟/OpenRouter 专项深析(读 review_json/focus_{key}.json)。"""
    ws = wb.create_sheet("重点分析")
    ws.append(["平台", "分析项", "分析内容"])
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN
    filled = False
    for key, title in FOCUS_TARGETS:
        p = os.path.join(BASE, "review_json", f"focus_{key}.json")
        if not os.path.exists(p):
            ws.append([title, "（待补充）", f"未找到 {p}，重点分析完成后重跑本脚本"])
            continue
        filled = True
        data = json.load(open(p, encoding="utf-8"))
        for item, content in data.items():
            ws.append([title, item, content])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
        for cell in row:
            cell.border = THIN
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 110
    if not filled:
        print("提示: 重点分析 sheet 为占位（focus_*.json 尚未生成）")


def sheet_training(wb):
    """模型层训练环节专项(读 review_json/training_focus.json)。"""
    ws = wb.create_sheet("模型层训练专项")
    head = ["厂商", "预训练语料中个人信息的告知与来源", "用户输入是否用于训练及opt-out",
            "To-B微调数据处理约定", "综合评述"]
    ws.append(head)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN
    p = os.path.join(BASE, "review_json", "training_focus.json")
    if not os.path.exists(p):
        ws.append(["（待补充）", f"未找到 {p}", "训练专项分析完成后重跑本脚本", "", ""])
    else:
        data = json.load(open(p, encoding="utf-8"))
        for vendor, items in data.items():
            ws.append([vendor] + [items.get(h, "") for h in head[1:]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5):
        for cell in row:
            cell.border = THIN
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 14
    for col in "BCDE":
        ws.column_dimensions[col].width = 45


def main():
    baseline, reviews, missing = load_all()
    if missing:
        raise SystemExit(f"缺少 {len(missing)} 个审查结果: {missing}")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_main(wb, baseline, reviews)
    sheet_summary(wb, baseline, reviews)
    sheet_prevalence(wb, baseline, reviews)
    sheet_focus(wb)
    sheet_training(wb)
    wb.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    main()
