# -*- coding: utf-8 -*-
"""Stage 1: Extract full text + tables from 7 brokerage contract .docx files."""

import json
import os
import sys
from pathlib import Path
import config

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

CONTRACTS_DIR = Path(config.BASE) / "经纪约word"
OUTPUT_DIR = Path(config.BASE) / "_brokerage_extract" / "_contracts_json"

# Real contracts (skip ~$ temp lock files)
CONTRACT_FILES = [
    "小鱼海棠21年签的经纪约（徐翔宇）—水火土(1).docx",
    "小鱼经纪约.docx",
    "阿飞-20年.docx",
    "北京车节奏-刘冠宇-新.docx",
    "大黄经纪约.docx",
    "良田经纪约.docx",
    "顾猛-车节奏最新pdf.docx",
]


def extract_contract(filepath: Path) -> dict:
    """Extract paragraphs, tables, and rels from a .docx contract."""
    doc = Document(str(filepath))

    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:  # skip empty paragraphs
            paragraphs.append({
                "idx": i,
                "text": text,
                "style": p.style.name if p.style else "None",
            })

    tables = []
    for t_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
        tables.append({"idx": t_idx, "rows": rows})

    # Detect embedded attachments / OLE objects
    rels_info = []
    try:
        for rel_id, rel in doc.part.rels.items():
            rels_info.append({
                "rel_id": rel_id,
                "rel_type": str(rel.reltype),
                "target_ref": str(rel.target_ref) if rel.target_ref else None,
            })
    except Exception as e:
        rels_info.append({"error": str(e)})

    # Build flat text for easy reading
    full_text_parts = []
    for p in paragraphs:
        full_text_parts.append(p["text"])
    for t in tables:
        for row in t["rows"]:
            full_text_parts.append(" | ".join(row))

    result = {
        "contract_name": filepath.name,
        "contract_stem": filepath.stem,
        "source_path": str(filepath),
        "paragraphs": paragraphs,
        "tables": tables,
        "rels": rels_info,
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "total_chars": sum(len(p["text"]) for p in paragraphs),
        "has_tables": len(tables) > 0,
    }

    return result


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for fname in CONTRACT_FILES:
        fpath = CONTRACTS_DIR / fname
        if not fpath.exists():
            print(f"  [SKIP] File not found: {fpath}")
            continue

        print(f"  Extracting: {fname} ...")
        try:
            data = extract_contract(fpath)
            out_path = OUTPUT_DIR / f"{fpath.stem}.json"
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            print(f"    -> {out_path.name}  ({data['paragraph_count']} paragraphs, {data['table_count']} tables, {data['total_chars']} chars)")
        except Exception as e:
            print(f"    [ERROR] {e}")

    print("\nDone. JSON files saved to:", str(OUTPUT_DIR))


if __name__ == "__main__":
    main()
